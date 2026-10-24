"""Streaming adapters for exported files; never execute HTML or embedded scripts."""
import codecs
import json
import re
from datetime import datetime, timezone, timedelta
from pathlib import Path

import ijson
import regex
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from echotrace.util import AppError

TZ = timezone(timedelta(hours=8))
ALIASES = {
    "ts": ("ts", "timestamp", "time", "msgTime", "时间", "发送时间", "消息时间"),
    "text": ("text", "content", "message", "内容", "消息内容", "文本"),
    "sender": ("sender", "senderName", "senderNick", "nickname", "发送人", "发送者", "昵称"),
    "msg_id": ("msg_id", "msgId", "message_id", "id", "消息ID"),
}


def field(row, name, default=""):
    return next((row[k] for k in ALIASES[name] if k in row and row[k] is not None), default)


def timestamp(value):
    if isinstance(value, datetime):
        return int(value.replace(tzinfo=value.tzinfo or TZ).timestamp())
    try:
        n = float(value)
        return int(n / 1000 if n > 100_000_000_000 else n)
    except (ValueError, TypeError):
        pass
    value = str(value).strip().replace("/", "-")
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return int(dt.replace(tzinfo=dt.tzinfo or TZ).timestamp())
    except ValueError as exc:
        raise AppError("IMPORT_PARSE_FAILED", f"无法识别时间：{value[:60]}。请使用日期时间或 Unix 秒。") from exc


def normalize(row):
    if not isinstance(row, dict):
        raise AppError("IMPORT_PARSE_FAILED", "每条消息必须是 JSON 对象。")
    if row.get("chatType", row.get("chat_type", 1)) not in (1, "1", "private", None):
        raise AppError("PRIVATE_CHAT_ONLY", "当前只支持私聊，请按联系人分别导出。")
    speaker = field(row, "sender")
    if isinstance(speaker, dict):
        name = speaker.get("name") or speaker.get("nickname") or speaker.get("nick") or speaker.get("uin") or speaker.get("uid")
        identity = speaker.get("uid") or speaker.get("uin") or name
    else:
        name = speaker
        identity = row.get("senderUid") or row.get("senderUin") or row.get("sender_id") or name
    name = str(name or identity or "未知发言人").strip().lstrip("@")
    content = field(row, "text")
    if isinstance(content, (dict, list)):
        content = json.dumps(content, ensure_ascii=False)
    kinds, parts, reply = [], [], row.get("reply_to", "")
    for element in row.get("elements", []) or []:
        typ = element.get("elementType", element.get("type"))
        if typ in (1, "text") or "textElement" in element:
            parts.append(str((element.get("textElement") or {}).get("content", element.get("text", ""))))
            kinds.append("text")
        elif typ in (2, "image") or "picElement" in element:
            parts.append("[图片]")
            kinds.append("image")
        elif typ in (6, "emoji", "face") or "faceElement" in element:
            face = element.get("faceElement") or {}
            parts.append(f"[表情:{face.get('faceText') or face.get('faceName') or face.get('faceIndex') or '表情'}]")
            kinds.append("emoji")
        elif typ in (3, "file"):
            parts.append("[文件]")
            kinds.append("file")
        elif typ in (4, "voice", "ptt"):
            parts.append("[语音]")
            kinds.append("voice")
        elif typ in (5, "video"):
            parts.append("[视频]")
            kinds.append("video")
        elif typ in (7, "reply") or "replyElement" in element:
            reply = (element.get("replyElement") or {}).get("replayMsgId", element.get("replyMsgId", ""))
            parts.append(f"[引用:{reply}]")
            kinds.append("reply")
        else:
            parts.append("[结构化消息]")
            kinds.append("json")
    text = "".join(parts) if parts else str(content or "")
    kind = row.get("kind") or ("text" if "text" in kinds else next(iter(kinds), "text"))
    if "撤回了一条消息" in text or row.get("recalled"):
        kind = "recall"
    if kind not in {"text", "image", "emoji", "voice", "video", "file", "recall", "system", "json", "reply"}:
        kind = "json"
    return {"msg_id": str(field(row, "msg_id")), "sender_raw": name, "identity": str(identity or name),
            "role_hint": row.get("sender") if row.get("sender") in ("me", "them", "system") else None,
            "ts": timestamp(field(row, "ts")), "text": text, "kind": kind, "reply_to": str(reply), "raw": row,
            "encoding": row.get("encoding"), "source_line": row.get("source_line")}


def json_rows(path):
    prefix = None
    with path.open("rb") as handle:
        for key, event, value in ijson.parse(handle):
            if event == "start_map" and key in ("item", "messages.item", "data.messages.item", "data.item", "records.item", "chatHistory.item"):
                prefix = key
                break
    if prefix is None:
        raise AppError("IMPORT_PARSE_FAILED", "JSON 中没有找到消息列表（messages / data.messages / 数组）。")
    with path.open("rb") as handle:
        yield from ijson.items(handle, prefix, use_float=True)


TEXT_PATTERN = r"^\[?(?P<ts>\d{4}[-/]\d{1,2}[-/]\d{1,2}[ T]\d{1,2}:\d{2}(?::\d{2})?)\]?\s+(?P<sender>[^:：\n]+)[:：][ \t]?(?P<text>.*)$"
HEADER_PATTERN = r"^(?P<ts>\d{4}[-/]\d{1,2}[-/]\d{1,2}[ T]\d{1,2}:\d{2}(?::\d{2})?)\s+(?P<sender>[^\n]+)$"


def _valid_text_headers(text, custom_pattern=None):
    if custom_pattern:
        try:
            pattern = regex.compile(custom_pattern)
        except regex.error as exc:
            raise AppError("INVALID_PATTERN", "自定义正则无效。") from exc
        return sum(bool(pattern.match(line, timeout=0.1)) for line in text.splitlines())
    return sum(bool(re.match(TEXT_PATTERN, line) or re.match(HEADER_PATTERN, line)) for line in text.splitlines())


def decode_qq_txt(path: Path, encoding=None, custom_pattern=None):
    data = path.read_bytes()
    requested = encoding.lower().replace("_", "-") if encoding else None
    bom = None
    if data.startswith(codecs.BOM_UTF8):
        bom = "utf-8-sig"
    elif data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        bom = "utf-16"
    if requested:
        selected = {"utf8": "utf-8", "utf16": "utf-16"}.get(requested, requested)
        if selected not in {"utf-8", "utf-8-sig", "gbk", "utf-16", "utf-16-le", "utf-16-be"}:
            raise AppError("IMPORT_ENCODING_INVALID", f"不支持的显式编码：{encoding}。")
        try:
            text = data.decode(selected)
        except UnicodeDecodeError as exc:
            raise AppError("IMPORT_ENCODING_FAILED", f"无法按显式编码 {encoding} 解码文件。") from exc
        if not _valid_text_headers(text, custom_pattern):
            raise AppError("IMPORT_ENCODING_FAILED", f"按显式编码 {encoding} 解码后未找到有效消息头。")
        return text, selected
    if bom:
        try:
            text = data.decode(bom)
        except UnicodeDecodeError as exc:
            raise AppError("IMPORT_ENCODING_FAILED", f"文件包含 {bom} BOM，但无法严格解码。") from exc
        if not _valid_text_headers(text, custom_pattern):
            raise AppError("IMPORT_PARSE_FAILED", "BOM 解码成功，但没有识别到有效消息头。")
        return text, bom
    candidates = []
    for candidate in ("utf-8", "gbk", "utf-16"):
        try:
            text = data.decode(candidate)
        except (UnicodeDecodeError, UnicodeError):
            continue
        if _valid_text_headers(text, custom_pattern):
            candidates.append((candidate, text))
    unique = {}
    for candidate, text in candidates:
        unique.setdefault(text, candidate)
    if not unique:
        raise AppError("IMPORT_ENCODING_FAILED", "无法用 UTF-8、GBK 或 UTF-16 严格解码并识别消息头，请显式选择编码。")
    if len(unique) > 1:
        raise AppError("IMPORT_ENCODING_AMBIGUOUS", "文件编码存在歧义，请显式选择 --encoding。")
    text, selected = next(iter(unique.items()))
    return text, selected


def text_rows(path, custom_pattern=None, encoding=None):
    try:
        pattern = regex.compile(custom_pattern or TEXT_PATTERN)
    except regex.error as exc:
        raise AppError("INVALID_PATTERN", "自定义正则无效。") from exc
    if not {"ts", "sender", "text"}.issubset(pattern.groupindex):
        raise AppError("INVALID_PATTERN", "正则必须含命名组 ts、sender、text。")
    pending, body_lines, matched = None, [], 0
    text, selected_encoding = decode_qq_txt(path, encoding, custom_pattern)

    def completed(separator=False):
        lines = body_lines[:-1] if separator and body_lines and body_lines[-1] == "" else body_lines
        item = dict(pending)
        inline = item.pop("text", "")
        item["text"] = "\n".join(([inline] if inline else []) + lines)
        return item

    for line_number, line in enumerate(text.splitlines(), 1):
        try:
            match = pattern.match(line, timeout=0.1)
        except TimeoutError as exc:
            raise AppError("INVALID_PATTERN", "正则匹配超时，请简化表达式。") from exc
        header = None if custom_pattern or match else re.match(HEADER_PATTERN, line)
        if match or header:
            if pending:
                yield completed(separator=True)
            pending = (match or header).groupdict()
            pending.setdefault("text", "")
            pending["encoding"] = selected_encoding
            pending["source_line"] = line_number
            body_lines = []
            matched += 1
        elif pending:
            body_lines.append(line)
        elif line.strip():
            raise AppError("IMPORT_PARSE_FAILED", f"第 {line_number} 行无法归属到任何消息，请核查原始记录。")
    if pending:
        yield completed()
    if not matched:
        raise AppError("IMPORT_PARSE_FAILED", "没有识别到消息，请使用「2026-01-01 12:00:00 昵称: 内容」格式或自定义正则。")


def html_rows(path):
    if path.stat().st_size > 80_000_000:
        raise AppError("IMPORT_PARSE_FAILED", "大型 HTML 请先从 QCE 导出 JSONL，避免占用过多内存。")
    soup = BeautifulSoup(path.read_text(encoding="utf-8-sig"), "html.parser")
    for script in soup.find_all("script"):
        text = script.string or script.get_text()
        candidates = [text] if script.get("type") == "application/json" else []
        match = re.search(r"(?:chatData|messages|exportData)\s*=\s*", text)
        if match:
            candidates.append(text[match.end():])
        for candidate in candidates:
            try:
                obj, _ = json.JSONDecoder().raw_decode(candidate.lstrip())
                rows = obj if isinstance(obj, list) else obj.get("messages", obj.get("data", {}).get("messages", []))
                if rows:
                    yield from rows
                    return
            except (ValueError, AttributeError):
                continue
    count = 0
    for el in soup.select(".message, .message-item, [data-message-id]"):
        name = el.select_one(".sender, .sender-name, .nickname, .message-sender")
        ts = el.select_one("time, .time, .timestamp, .message-time")
        body = el.select_one(".content, .message-content, .text")
        if (ts or el.get("data-timestamp")) and name:
            count += 1
            yield {"msg_id": el.get("data-message-id", ""), "sender": name.get_text(strip=True),
                   "ts": el.get("data-timestamp") or ts.get("datetime") or ts.get_text(strip=True),
                   "text": body.get_text("\n", strip=True) if body else "[图片]"}
    if not count:
        raise AppError("IMPORT_PARSE_FAILED", "此 HTML 布局暂不支持，请改用 JSON 或 JSONL 导出。")


def parse_file(path: Path, custom_pattern=None, encoding=None):
    ext = path.suffix.lower()
    try:
        if ext in (".jsonl", ".ndjson"):
            with path.open(encoding="utf-8-sig") as handle:
                for line_number, line in enumerate(handle, 1):
                    if line.strip():
                        row = json.loads(line)
                        if row.get("type") in ("metadata", "header", "footer"):
                            continue
                        yield normalize(row.get("message", row) if isinstance(row.get("message"), dict) else row)
        elif ext == ".json":
            for row in json_rows(path):
                yield normalize(row)
        elif ext in (".txt", ".log"):
            for row in text_rows(path, custom_pattern, encoding):
                yield normalize(row)
        elif ext in (".html", ".htm"):
            for row in html_rows(path):
                yield normalize(row)
        elif ext in (".xlsx", ".xlsm"):
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    rows = sheet.iter_rows(values_only=True)
                    keys = [str(v).strip() for v in next(rows, ())]
                    for row in rows:
                        if any(v is not None for v in row):
                            yield normalize(dict(zip(keys, row)))
            finally:
                book.close()
        else:
            raise AppError("IMPORT_PARSE_FAILED", "支持 JSON / JSONL / TXT / HTML / XLSX；旧版 XLS 请另存为 XLSX。")
    except AppError:
        raise
    except Exception as exc:
        raise AppError("IMPORT_PARSE_FAILED", f"文件解析失败：{type(exc).__name__}。请检查编码、字段与导出格式。") from exc
